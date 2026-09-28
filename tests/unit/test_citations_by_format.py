"""A citation sends the teacher somewhere she can look, whatever the format.

Measured on 2026-09-25 with ``scripts/eval/check_formats.py``, on Word and
PowerPoint files built from real pages of the ministerial curriculum:

* a slide was cited as "p. 3" -- a page no slideshow has;
* a Word passage was cited as "(file.docx)" with no location at all, although
  its chunk already carried the section heading;
* the speaker notes of a deck were not read: 4 of 4 sentences that lived only
  in the notes were missing from the index.

These tests go from real files, built in memory, to the sentence the teacher
reads -- parser, chunker, marker, renderer -- because each defect lived in the
hand-off between two of them.
"""

from __future__ import annotations

import io
import uuid

import pytest

from aix.webui.materials.chunking import chunk_document
from aix.webui.materials.citations import render_citations, sources_block
from aix.webui.materials.parsing import check_upload, parse_document
from aix.webui.materials.retrieval import MaterialContextItem

pytestmark = pytest.mark.unit


def _deck() -> bytes:
    from pptx import Presentation

    deck = Presentation()
    for title, body, notes in [
        ("Geografia", "La geografia studia i rapporti fra le società e il pianeta.", "Nota: partire da una mappa del quartiere."),
        ("Storia", "La storia si apre al confronto fra fonti diverse.", ""),
    ]:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text_frame.text = body
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def _word() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("Musica", level=1)
    doc.add_paragraph("L'apprendimento della musica consta di pratiche e di conoscenze.")
    doc.add_heading("Arte & immagine = linguaggi [visivi]", level=1)
    doc.add_paragraph("La disciplina sviluppa la capacità di leggere le immagini.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _chunks(name: str, content: bytes):
    accepted = check_upload(name, content)
    return chunk_document(parse_document(content, accepted.media_type))


def _cite(chunk, name: str) -> tuple[str, str]:
    """The marker the Writer would copy, and what the teacher would read."""
    source = str(uuid.uuid4())
    item = MaterialContextItem(source_id=source, text=chunk.text, score=1.0, locator=dict(chunk.locator))
    return item.citation, render_citations(item.citation, {source: name})


def _containing(chunks, phrase: str):
    return next(c for c in chunks if phrase in c.text)


# ---------------------------------------------------------------------------
# PowerPoint
# ---------------------------------------------------------------------------


def test_a_slide_is_cited_as_a_slide_not_a_page():
    chunks = _chunks("lezione.pptx", _deck())

    _marker, read = _cite(_containing(chunks, "confronto fra fonti"), "lezione.pptx")

    assert read == "(lezione.pptx, slide 2)"


def test_speaker_notes_are_read_and_belong_to_their_slide():
    chunks = _chunks("lezione.pptx", _deck())

    chunk = _containing(chunks, "mappa del quartiere")

    assert chunk.locator["page"] == 1
    assert _cite(chunk, "lezione.pptx")[1] == "(lezione.pptx, slide 1)"


def test_every_chunk_says_which_format_it_came_from():
    assert {c.locator.get("format") for c in _chunks("lezione.pptx", _deck())} == {"pptx"}
    assert {c.locator.get("format") for c in _chunks("appunti.docx", _word())} == {"docx"}


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------


def test_a_word_passage_is_cited_by_its_section():
    chunks = _chunks("appunti.docx", _word())

    _marker, read = _cite(_containing(chunks, "pratiche e di conoscenze"), "appunti.docx")

    assert read == "(appunti.docx, § Musica)"


def test_a_heading_with_marker_syntax_in_it_survives_the_marker():
    """``&``, ``=`` and ``]`` are the marker's own punctuation."""
    chunks = _chunks("appunti.docx", _word())

    marker, read = _cite(_containing(chunks, "leggere le immagini"), "appunti.docx")

    assert marker.count("]]") == 1, "the heading must not close the marker early"
    assert read == "(appunti.docx, § Arte & immagine = linguaggi [visivi])"


# ---------------------------------------------------------------------------
# Everything else, unchanged or simpler
# ---------------------------------------------------------------------------


def test_a_pdf_page_is_still_a_page():
    item = MaterialContextItem(
        source_id="11111111-1111-1111-1111-111111111111",
        text="x",
        score=1.0,
        locator={"format": "pdf", "page": 7, "block_ids": ["b1"]},
    )

    assert render_citations(item.citation, {item.source_id: "guida.pdf"}) == "(guida.pdf, p. 7)"


def test_a_chunk_indexed_before_the_format_was_recorded_reads_as_before():
    """Old chunks have no ``format``: their page stays a page, nothing breaks."""
    item = MaterialContextItem(
        source_id="11111111-1111-1111-1111-111111111111", text="x", score=1.0, locator={"page": 3}
    )

    assert render_citations(item.citation, {item.source_id: "vecchio.pdf"}) == "(vecchio.pdf, p. 3)"


def test_a_photo_is_not_given_a_page_number():
    item = MaterialContextItem(
        source_id="11111111-1111-1111-1111-111111111111",
        text="x",
        score=1.0,
        locator={"format": "image", "page": 1, "block_ids": ["b1"], "ocr": True},
    )

    assert render_citations(item.citation, {item.source_id: "lavagna.jpg"}) == "(lavagna.jpg)"


def test_the_sources_list_speaks_each_format():
    deck, word = "22222222-2222-2222-2222-222222222222", "33333333-3333-3333-3333-333333333333"
    text = (
        f"Uno [[material:{deck}#slide=4&block=a]]. Due [[material:{deck}#slide=2&block=b]]. "
        f"Tre [[material:{word}#section=Storia&block=c]]. Quattro [[material:{word}#section=Musica&block=d]]."
    )

    block = sources_block(text, {deck: "lezione.pptx", word: "appunti.docx"})

    assert "- lezione.pptx — slide 2, 4" in block
    assert "- appunti.docx — § Storia, § Musica" in block, "sections in the order first cited"
