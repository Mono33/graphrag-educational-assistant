"""Pictures inside PowerPoint and Word files are read, like uploaded photos.

A slide that is a map, a chart or a screenshot holds its text in the picture;
the parser used to skip it. OCR is replaced by a fake reader: nothing leaves
the machine.
"""

from __future__ import annotations

import io

import pytest
from PIL import Image

from aix.webui.materials.parsing import parse_document, parsers
from aix.webui.materials.parsing.intake import DOCX, PPTX
from aix.webui.materials.parsing.ocr import OcrResult

pytestmark = pytest.mark.unit

MAP_TEXT = "La Guerra Fredda 1962: basi aeree statunitensi e missili balistici"


def _png(width: int, height: int, colour: str) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (width, height), colour).save(buffer, format="PNG")
    return buffer.getvalue()


MAP = _png(800, 500, "orange")
PHOTO = _png(600, 400, "grey")
LOGO = _png(300, 150, "navy")
ICON = _png(40, 40, "red")


@pytest.fixture
def reader(monkeypatch):
    """The map has text, the photo has none, the logo has a word or two."""
    calls: list[tuple[int, int]] = []

    def _recognise(image, timeout_s):
        calls.append(image.size)
        text = {
            (800, 500): MAP_TEXT,
            (600, 400): "[figura]",
            (300, 150): "Liceo Galilei — dipartimento di storia",
        }[image.size]
        return OcrResult(text, 0.95, len(text.split()), engine="vision")

    monkeypatch.setattr(parsers, "ocr_available", lambda: True)
    monkeypatch.setattr(parsers, "recognise", _recognise)
    return calls


def _deck() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    blank = deck.slide_layouts[6]
    one = deck.slides.add_slide(blank)
    one.shapes.add_picture(io.BytesIO(MAP), Inches(1), Inches(1))
    one.shapes.add_picture(io.BytesIO(ICON), Inches(0.1), Inches(0.1))
    one.shapes.add_picture(io.BytesIO(LOGO), Inches(6), Inches(6))
    two = deck.slides.add_slide(blank)
    two.shapes.add_picture(io.BytesIO(PHOTO), Inches(1), Inches(1))
    two.shapes.add_picture(io.BytesIO(LOGO), Inches(6), Inches(6))
    group = two.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1))
    box.text_frame.text = "Didascalia: Yalta, febbraio 1945"
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def test_the_text_in_a_slide_s_picture_is_read_and_cited_by_slide(reader):
    ir = parse_document(_deck(), PPTX)

    block = next(b for b in ir.blocks if MAP_TEXT in b.text)
    assert block.page == 1, "cited as slide 1"
    assert block.source_locator["image"] is True
    assert block.source_locator["ocr_engine"] == "vision"
    assert block.confidence < 1.0, "marked as OCR text downstream"
    assert ir.ocr_pages == (1,)


def test_a_lone_figure_marker_is_not_a_paragraph(reader, monkeypatch):
    """The real deck came back as «[figura]» + «MRBM FIELD LAUNCH SITE…»."""
    monkeypatch.setattr(
        parsers,
        "recognise",
        lambda image, timeout_s: OcrResult(f"[figura]\n\n{MAP_TEXT}", 0.95, 12, engine="vision"),
    )

    ir = parse_document(_deck(), PPTX)

    assert not any(b.text.strip() == "[figura]" for b in ir.blocks)
    assert any(MAP_TEXT in b.text for b in ir.blocks)


def test_icons_are_skipped_and_a_repeated_logo_is_read_once(reader):
    ir = parse_document(_deck(), PPTX)

    assert (40, 40) not in reader, "an icon is never sent"
    assert reader.count((300, 150)) == 1, "the logo on both slides is read once"
    assert sum("Liceo Galilei" in b.text for b in ir.blocks) == 1


def test_a_photo_without_text_is_neither_indexed_nor_a_warning(reader):
    ir = parse_document(_deck(), PPTX)

    assert not any("[figura]" in b.text for b in ir.blocks)
    assert not any("page=2" in w for w in ir.warnings)


def test_text_inside_a_group_of_shapes_is_read(reader):
    ir = parse_document(_deck(), PPTX)

    assert any("Yalta, febbraio 1945" in b.text and b.page == 2 for b in ir.blocks)


def test_the_page_cap_applies_to_pictures_too(reader):
    ir = parse_document(_deck(), PPTX, parsers.ParseOptions(max_ocr_pages=1))

    assert len(reader) == 1
    assert "ocr_page_limit:1" in ir.warnings


def test_a_picture_in_a_word_file_is_read_after_its_paragraph(reader):
    import docx
    from docx.shared import Inches

    document = docx.Document()
    document.add_heading("La Guerra Fredda", level=1)
    document.add_paragraph("La mappa mostra le basi del 1962.")
    document.add_picture(io.BytesIO(MAP), width=Inches(4))
    document.add_paragraph("Dopo la mappa, il commento.")
    buffer = io.BytesIO()
    document.save(buffer)

    ir = parse_document(buffer.getvalue(), DOCX)

    texts = [b.text for b in ir.blocks]
    assert texts.index(MAP_TEXT) == texts.index("La mappa mostra le basi del 1962.") + 1
    assert texts.index("Dopo la mappa, il commento.") == texts.index(MAP_TEXT) + 1
    assert ir.parser_version == "2"


def test_without_ocr_the_pictures_are_left_as_before(monkeypatch):
    monkeypatch.setattr(parsers, "ocr_available", lambda: False)

    ir = parse_document(_deck(), PPTX)

    assert not any(MAP_TEXT in b.text for b in ir.blocks)
    assert any("Yalta" in b.text for b in ir.blocks), "the grouped text needs no OCR"
