"""The Word export, and the route that serves it (2026-10-14).

Checked by reading the file back with python-docx -- its text, its styles, its
links, its header and footer -- as the PDF is read back with pypdf: what a
teacher needs is that the words are there, the structure is Word's own and the
disclosure is on the page.
"""

from __future__ import annotations

import io

import pytest
from docx import Document
from tests.unit.test_export_pdf import (  # noqa: F401 — the fixture and the lesson
    BRUNO,
    LESSON,
    LESSON_ID,
    _export,
    client,
)

from aix.export import adapt_lesson_markdown, build_export_ir
from aix.export.word import render_docx

pytestmark = pytest.mark.unit

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _docx(markdown: str = LESSON):
    content = render_docx(build_export_ir(adapt_lesson_markdown(markdown, lesson_id="L1"))).content
    return Document(io.BytesIO(content))


@pytest.fixture(scope="module")
def doc():
    return _docx()


def _paragraphs(doc):
    return [(p.style.name, p.text) for p in doc.paragraphs]


def test_the_lesson_text_is_in_the_file_with_its_characters(doc):
    text = "\n".join(p.text for p in doc.paragraphs)

    assert "Memoria di lavoro — perché dimentichiamo" in text
    # Word draws what the PDF font cannot: nothing is substituted or dropped.
    assert "una freccia → con ≤ e ≥" in text
    assert "🧠" in text


def test_headings_are_word_headings(doc):
    paragraphs = _paragraphs(doc)

    assert ("Title", "Memoria di lavoro — perché dimentichiamo") in paragraphs
    assert ("Heading 1", "Introduzione (10 minuti)") in paragraphs
    assert ("Heading 2", "Fase 1 — I Do (25 minuti)") in paragraphs


def test_lists_keep_their_nesting_and_their_own_numbers(doc):
    paragraphs = _paragraphs(doc)

    assert ("List Bullet", "Spiegazione con schede cartacee") in paragraphs
    assert ("List Bullet 2", "una scheda per coppia") in paragraphs
    texts = [text for _, text in paragraphs]
    assert "1.\tLavoro in gruppi da 4" in texts
    assert "2.\tRestituzione" in texts


def test_a_table_is_a_word_table_with_its_header_repeated(doc):
    table = doc.tables[0]

    assert [c.text for c in table.rows[0].cells] == ["Strumento", "Funzione"]
    assert [c.text for c in table.rows[1].cells] == ["Exit ticket", "Verifica immediata"]
    assert table.rows[0]._tr.xpath("./w:trPr/w:tblHeader")


def test_links_are_clickable(doc):
    targets = [r.target_ref for r in doc.part.rels.values() if r.reltype.endswith("/hyperlink")]

    assert targets == ["https://youtu.be/bSycdIx-C48"]


def test_an_unsafe_link_is_text_only():
    doc = _docx("# T\n\nUn [link](javascript:alert(1)) e uno [buono](https://fem.digital).")
    targets = [r.target_ref for r in doc.part.rels.values() if r.reltype.endswith("/hyperlink")]

    assert targets == ["https://fem.digital"], "the javascript: link is never made clickable"


def test_the_ai_disclosure_is_on_every_page_and_in_the_properties(doc):
    section = doc.sections[0]

    assert "IA" in section.header.paragraphs[0].text
    assert "intelligenza artificiale" in section.footer.paragraphs[0].text
    assert "PAGE" in section.footer.paragraphs[0]._p.xml
    assert "intelligenza artificiale" in doc.core_properties.comments
    assert doc.core_properties.subject.startswith("lesson=L1 content=")


def test_control_characters_do_not_break_the_file():
    doc = _docx("# Titolo\x0b\n\nTesto con un carattere\x01 di controllo.")

    assert "Testo con un carattere di controllo." in "\n".join(p.text for p in doc.paragraphs)


def test_an_empty_lesson_still_renders():
    assert _docx("").paragraphs


def test_the_route_serves_a_word_download(client):
    response = _export(client, "docx")

    assert response.status_code == 200
    assert response.headers["content-type"] == DOCX_TYPE
    assert response.headers["content-disposition"].endswith('.docx"')
    assert response.headers["x-ai-generated"] == "true"
    text = "\n".join(p.text for p in Document(io.BytesIO(response.content)).paragraphs)
    assert "Memoria di lavoro" in text


def test_switched_off_the_route_refuses_it(client, monkeypatch):
    monkeypatch.setenv("AIX_EXPORT_DOCX_ENABLED", "false")

    assert _export(client, "docx").status_code == 404


def test_another_teacher_cannot_export_the_lesson(client):
    client.act_as(BRUNO)

    assert _export(client, "docx").status_code == 404
