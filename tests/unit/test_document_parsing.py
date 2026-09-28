"""Intake and parsing of uploaded material (Fase 2 · PR 5).

Fixtures are built in the test rather than stored on disk: a DOCX with python-docx,
a PPTX with python-pptx, a digital PDF with ReportLab, a scanned PDF by saving an
image as PDF. That way each test states exactly what the file contains, and the
parser is checked against a known truth rather than against a sample nobody can
describe.

The OCR tests need the Tesseract binary. They skip where it is absent — this
Windows machine — and run in the Linux image, where apt installs it.
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest

from aix.webui.materials.parsing import IntakeRejected, check_upload, parse_document
from aix.webui.materials.parsing.intake import DOCX, MARKDOWN, PDF, PPTX, TEXT
from aix.webui.materials.parsing.ocr import tesseract_available

pytestmark = pytest.mark.unit

FONT = (
    Path(__file__).resolve().parents[2]
    / "src"
    / "aix"
    / "export"
    / "fonts"
    / "NotoSans-Regular.ttf"
)
needs_tesseract = pytest.mark.skipif(not tesseract_available(), reason="tesseract not installed")


# ---------------------------------------------------------------------------
# Fixture builders
# ---------------------------------------------------------------------------


def _digital_pdf(pages: list[str]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    for text in pages:
        y = 780
        for line in text.split("\n"):
            pdf.drawString(72, y, line)
            y -= 18
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def _text_image(text: str, size=(1600, 500)):
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(str(FONT), 64)
    y = 60
    for line in text.split("\n"):
        draw.text((60, y), line, fill="black", font=font)
        y += 110
    return image


def _png(image) -> bytes:
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def _scanned_pdf(image) -> bytes:
    buffer = io.BytesIO()
    image.save(buffer, format="PDF", resolution=200)
    return buffer.getvalue()


def _docx(with_macro: bool = False) -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("Programma di classe", level=1)
    document.add_paragraph("Si affronta la memoria di lavoro.")
    document.add_paragraph("Primo obiettivo", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Strumento", "Funzione"
    table.cell(1, 0).text, table.cell(1, 1).text = "Exit ticket", "Verifica"
    document.add_heading("Valutazione", level=2)

    buffer = io.BytesIO()
    document.save(buffer)
    content = buffer.getvalue()

    if with_macro:
        out = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(content)) as src, zipfile.ZipFile(out, "w") as dst:
            for item in src.infolist():
                dst.writestr(item, src.read(item.filename))
            dst.writestr("word/vbaProject.bin", b"\x00" * 64)
        content = out.getvalue()
    return content


def _pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = "La memoria di lavoro"

    bullets = presentation.slides.add_slide(presentation.slide_layouts[1])
    bullets.shapes.title.text = "Obiettivi"
    frame = bullets.placeholders[1].text_frame
    frame.text = "Descrivere le fasi"
    sub = frame.add_paragraph()
    sub.text = "encoding e retrieval"
    sub.level = 1

    table_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    table_slide.shapes.title.text = "Valutazione"
    table = table_slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Strumento", "Funzione"
    table.cell(1, 0).text, table.cell(1, 1).text = "Quiz", "Formativa"

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Intake: what may be parsed at all
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "name, builder, expected",
    [
        ("programma.pdf", lambda: _digital_pdf(["Testo"]), PDF),
        ("programma.docx", _docx, DOCX),
        ("lezione.pptx", _pptx, PPTX),
        ("note.txt", lambda: "Appunti in italiano: è così.".encode(), TEXT),
        ("note.md", lambda: b"# Titolo\n\nTesto.", MARKDOWN),
        ("scheda.png", lambda: _png(_text_image("Scheda")), "image/png"),
    ],
)
def test_the_allowed_formats_are_accepted(name, builder, expected):
    assert check_upload(name, builder()).media_type == expected


@pytest.mark.parametrize(
    "name, content, code",
    [
        ("programma.pdf.exe", b"MZ\x90\x00", "active_extension"),
        ("pagina.html", b"<script>alert(1)</script>", "active_extension"),
        ("disegno.svg", b"<svg onload='x()'/>", "active_extension"),
        ("archivio.zip", b"PK\x03\x04", "active_extension"),
        ("foglio.xlsx", b"PK\x03\x04", "extension_not_allowed"),
        ("senza_estensione", b"%PDF-1.4", "no_extension"),
        ("vuoto.pdf", b"", "empty"),
    ],
)
def test_dangerous_or_unknown_names_are_refused(name, content, code):
    with pytest.raises(IntakeRejected) as refused:
        check_upload(name, content)
    assert refused.value.code == code


def test_a_pdf_renamed_docx_is_refused_not_misparsed():
    with pytest.raises(IntakeRejected) as refused:
        check_upload("programma.docx", _digital_pdf(["Testo"]))
    assert refused.value.code in ("not_ooxml", "magic_mismatch")


def test_a_png_renamed_pdf_is_refused():
    with pytest.raises(IntakeRejected) as refused:
        check_upload("programma.pdf", _png(_text_image("x")))
    assert refused.value.code == "magic_mismatch"


def test_a_macro_enabled_document_is_refused_even_when_renamed():
    with pytest.raises(IntakeRejected) as refused:
        check_upload("programma.docx", _docx(with_macro=True))
    assert refused.value.code == "macro_enabled"


def test_an_encrypted_pdf_is_refused():
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    for page in PdfReader(io.BytesIO(_digital_pdf(["Riservato"]))).pages:
        writer.add_page(page)
    writer.encrypt("password")
    buffer = io.BytesIO()
    writer.write(buffer)

    with pytest.raises(IntakeRejected) as refused:
        check_upload("riservato.pdf", buffer.getvalue())
    assert refused.value.code == "encrypted"


def test_a_legacy_or_encrypted_office_container_is_refused():
    with pytest.raises(IntakeRejected) as refused:
        check_upload("vecchio.docx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512)
    assert refused.value.code == "ole_container"


def test_an_archive_bomb_is_refused_before_extraction():
    """A DOCX is a ZIP, and a ZIP is how 10 MB becomes 10 GB in memory."""
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as bomb:
        bomb.writestr("[Content_Types].xml", "<Types/>")
        bomb.writestr("word/document.xml", b"\x00" * (30 * 1024 * 1024))

    with pytest.raises(IntakeRejected) as refused:
        check_upload("bomba.docx", out.getvalue())
    assert refused.value.code == "archive_bomb"


def test_a_pdf_protected_only_against_editing_is_accepted():
    """Like the DM 254/2012 in the corpus: encrypted, but with no opening password."""
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    for page in PdfReader(
        io.BytesIO(_digital_pdf(["Indicazioni nazionali per il curricolo"]))
    ).pages:
        writer.add_page(page)
    writer.encrypt(user_password="", owner_password="ministero")
    buffer = io.BytesIO()
    writer.write(buffer)

    assert check_upload("indicazioni.pdf", buffer.getvalue()).page_count == 1
    ir = parse_document(buffer.getvalue(), PDF)
    assert "Indicazioni nazionali" in ir.text


def test_a_pdf_over_the_page_limit_is_refused():
    with pytest.raises(IntakeRejected) as refused:
        check_upload("lungo.pdf", _digital_pdf(["p"] * 201))
    assert refused.value.code == "too_many_pages"


def test_ocr_stops_at_its_own_page_cap(monkeypatch):
    """A long scan is read up to the cap and the rest is declared, not dropped silently."""
    from aix.webui.materials.parsing import ParseOptions, parsers
    from aix.webui.materials.parsing.ocr import OcrResult

    calls = []
    monkeypatch.setattr(parsers, "ocr_available", lambda: True)
    monkeypatch.setattr(parsers, "render_pdf_page", lambda content, index: index)
    monkeypatch.setattr(
        parsers,
        "recognise",
        lambda image, timeout_s: calls.append(image) or OcrResult(f"pagina {image}", 0.9, 2),
    )
    from pypdf import PdfReader, PdfWriter

    scan = PdfReader(io.BytesIO(_scanned_pdf(_text_image("x")))).pages[0]
    writer = PdfWriter()
    for _ in range(5):
        writer.add_page(scan)
    buffer = io.BytesIO()
    writer.write(buffer)

    ir = parse_document(buffer.getvalue(), PDF, ParseOptions(max_ocr_pages=2))

    assert len(calls) == 2
    assert ir.ocr_pages == (1, 2)
    assert ir.unread_pages == (3, 4, 5)
    assert "ocr_page_limit:2" in ir.warnings


def test_binary_data_in_a_text_file_is_refused():
    with pytest.raises(IntakeRejected) as refused:
        check_upload("note.txt", b"testo\x00binario")
    assert refused.value.code == "magic_mismatch"


def test_a_corrupt_pdf_is_a_rejection_not_a_crash():
    """The legacy upload returned a 500 for this. Intake answers it."""
    with pytest.raises(IntakeRejected) as refused:
        check_upload("rotto.pdf", b"%PDF-1.4\n questo non e' un pdf")
    assert refused.value.code == "unreadable_pdf"


def test_rejections_speak_italian_to_the_teacher():
    with pytest.raises(IntakeRejected) as refused:
        check_upload("programma.pdf.exe", b"MZ")
    assert "non consentito" in str(refused.value)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def test_a_digital_pdf_is_read_page_by_page_with_locators():
    ir = parse_document(
        _digital_pdf(["Programma di classe per la memoria", "Seconda pagina con la valutazione"]),
        PDF,
    )

    assert ir.page_count == 2
    assert [b.page for b in ir.blocks] == [1, 2]
    assert ir.blocks[0].source_locator == {"format": "pdf", "page": 1}
    assert all(b.confidence == 1.0 for b in ir.blocks)
    assert ir.ocr_pages == ()


def test_a_page_without_text_is_unread_when_ocr_is_off():
    from aix.webui.materials.parsing import ParseOptions

    ir = parse_document(_scanned_pdf(_text_image("Scheda")), PDF, ParseOptions(ocr_enabled=False))

    assert ir.unread_pages == (1,)
    assert ir.blocks == ()


def test_quality_reports_pages_not_just_empty():
    """The corpus scan had text on 1 page of 50 and still read as not empty."""
    from aix.webui.materials.parsing import ParseOptions

    mixed = _digital_pdf(["Copertina con abbastanza testo da contare davvero"])
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    writer.add_page(PdfReader(io.BytesIO(mixed)).pages[0])
    for _ in range(3):
        writer.add_page(PdfReader(io.BytesIO(_scanned_pdf(_text_image("x")))).pages[0])
    buffer = io.BytesIO()
    writer.write(buffer)

    quality = parse_document(buffer.getvalue(), PDF, ParseOptions(ocr_enabled=False)).quality()

    assert quality["empty"] is False
    assert quality["pages"] == 4
    assert quality["pages_with_text_layer"] == 1
    assert quality["unread_pages"] == 3, "the signal `empty` alone hid"


# ---------------------------------------------------------------------------
# DOCX and PPTX
# ---------------------------------------------------------------------------


def test_a_docx_keeps_headings_lists_and_tables_in_order():
    ir = parse_document(_docx(), DOCX)

    assert [b.kind for b in ir.blocks] == ["heading", "paragraph", "list_item", "table", "heading"]
    assert ir.blocks[0].level == 1 and ir.blocks[-1].level == 2
    assert ir.blocks[3].rows == (("Strumento", "Funzione"), ("Exit ticket", "Verifica"))
    assert all(b.source_locator["format"] == "docx" for b in ir.blocks)


def test_a_pptx_is_read_slide_by_slide():
    ir = parse_document(_pptx(), PPTX)

    assert ir.page_count == 3
    titles = [b for b in ir.blocks if b.kind == "slide_title"]
    assert [t.text for t in titles] == ["La memoria di lavoro", "Obiettivi", "Valutazione"]
    nested = next(b for b in ir.blocks if b.text == "encoding e retrieval")
    assert nested.kind == "list_item"
    assert nested.source_locator["slide_index"] == 2
    table = next(b for b in ir.blocks if b.kind == "table")
    assert table.rows[1] == ("Quiz", "Formativa")


def test_markdown_headings_are_recognised():
    ir = parse_document(b"# Titolo\n\nUn paragrafo.\n\n## Sezione", MARKDOWN)

    assert [(b.kind, b.level) for b in ir.blocks] == [
        ("heading", 1),
        ("paragraph", None),
        ("heading", 2),
    ]


def test_block_ids_are_stable_for_the_same_bytes():
    content = _docx()
    assert [b.block_id for b in parse_document(content, DOCX).blocks] == [
        b.block_id for b in parse_document(content, DOCX).blocks
    ]


def test_a_disabled_format_is_refused_by_its_own_flag(monkeypatch):
    from aix.webui.materials.parsing.parsers import FormatDisabled

    monkeypatch.setenv("AIX_PARSE_PPTX_ENABLED", "false")
    with pytest.raises(FormatDisabled):
        parse_document(_pptx(), PPTX)
    assert parse_document(_docx(), DOCX).blocks, "the other formats are unaffected"


# ---------------------------------------------------------------------------
# OCR — run in the Linux image, skipped where Tesseract is absent
# ---------------------------------------------------------------------------


@needs_tesseract
def test_an_image_is_read_by_ocr_with_a_confidence():
    ir = parse_document(_png(_text_image("Programma di classe\nmemoria di lavoro")), "image/png")

    assert ir.ocr_pages == (1,)
    assert "Programma" in ir.text
    assert 0.0 < ir.blocks[0].confidence <= 1.0
    assert ir.blocks[0].source_locator["ocr"] is True


@needs_tesseract
def test_only_the_pages_without_a_text_layer_are_ocrd():
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    writer.add_page(
        PdfReader(io.BytesIO(_digital_pdf(["Pagina digitale con testo sufficiente"]))).pages[0]
    )
    writer.add_page(PdfReader(io.BytesIO(_scanned_pdf(_text_image("Scheda scansionata")))).pages[0])
    buffer = io.BytesIO()
    writer.write(buffer)

    ir = parse_document(buffer.getvalue(), PDF)

    assert ir.ocr_pages == (2,), "page 1 has a text layer and is never rasterised"
    assert "Scheda" in ir.text
    assert ir.quality()["pages_with_text_layer"] == 1
