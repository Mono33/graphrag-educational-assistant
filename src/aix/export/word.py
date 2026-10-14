"""The lesson as a Word document (2026-10-14).

Teachers adapt a lesson before using it: a PDF is for printing, a Word file is
for editing. This renders ``ExportIRV1`` with python-docx -- the same snapshot
the PDF reads, so the two carry the same sections, the same sources and the
same Art. 50 notice -- into real Word structure rather than formatted text:

* **headings** are Word headings (``##`` → Heading 1), so the navigation pane
  and a table of contents work;
* **lists** are bulleted paragraphs; a numbered list is numbered in its own
  text, because Word's numbered style continues across separate lists and a
  second list would start at 4;
* **tables** are Word tables with the header row repeated across pages;
* **links** are real hyperlinks, for the schemes the adapter allows;
* **the disclosure** is in the page header and footer of every page, with the
  page number, and in the document properties.

Unlike the PDF there is no glyph filter: Word substitutes a font for a
character the chosen one lacks, so an arrow or an emoji survives. What XML
cannot hold at all -- control characters -- is removed.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from aix.export.export_ir import ExportIRV1
from aix.export.lesson_document import Block, Run

RENDERER_VERSION = "docx@1"

_MUTED = RGBColor(0x6B, 0x72, 0x80)
_LINK = RGBColor(0x1F, 0x5F, 0xA8)
_CODE = RGBColor(0x4B, 0x55, 0x63)
_HEADER_FILL = "EEF2F6"

#: Characters XML 1.0 cannot carry: python-docx raises on them.
_XML_INVALID = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


@dataclass(frozen=True)
class DocxResult:
    content: bytes


def _clean(text: str) -> str:
    return _XML_INVALID.sub("", text or "")


# ---------------------------------------------------------------------------
# Inline
# ---------------------------------------------------------------------------


def _hyperlink(paragraph, run: Run, text: str) -> None:
    """A clickable link: python-docx has no API for it, so the XML is built here."""
    rel_id = paragraph.part.relate_to(run.href, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    r = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), str(_LINK))
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    if run.bold:
        props.append(OxmlElement("w:b"))
    if run.italic:
        props.append(OxmlElement("w:i"))
    r.append(props)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    link.append(r)
    paragraph._p.append(link)


def _runs(paragraph, runs: tuple[Run, ...]) -> None:
    for run in runs:
        text = _clean(run.text)
        if not text:
            continue
        if run.href:
            _hyperlink(paragraph, run, text)
            continue
        out = paragraph.add_run(text)
        out.bold = run.bold or None
        out.italic = run.italic or None
        if run.code:
            out.font.name = "Consolas"
            out.font.color.rgb = _CODE


# ---------------------------------------------------------------------------
# Blocks
# ---------------------------------------------------------------------------


def _heading(doc, text: str, level: int) -> None:
    # ``##`` is the first level inside the lesson: the title above it is "Title".
    doc.add_heading(_clean(text), level=min(max(level - 1, 1), 3))


def _list(doc, block: Block) -> None:
    counter = 0
    for runs, depth in zip(block.items, block.depths, strict=True):
        if block.ordered and depth == 0:
            counter += 1
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(18)
            paragraph.paragraph_format.first_line_indent = Pt(-14)
            paragraph.add_run(f"{counter}.\t")
        else:
            style = "List Bullet" if depth == 0 else f"List Bullet {min(depth + 1, 3)}"
            paragraph = doc.add_paragraph(style=style)
        _runs(paragraph, runs)


def _repeat_header(row) -> None:
    props = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    props.append(header)


def _shade(cell) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), _HEADER_FILL)
    cell._tc.get_or_add_tcPr().append(shading)


def _table(doc, block: Block) -> None:
    if not block.rows:
        return
    columns = max(len(row) for row in block.rows)
    table = doc.add_table(rows=len(block.rows), cols=columns)
    table.style = "Table Grid"
    for r, row in enumerate(block.rows):
        for c in range(columns):
            cell = table.cell(r, c)
            paragraph = cell.paragraphs[0]
            if c < len(row):
                _runs(paragraph, row[c])
            if r == 0:
                for out in paragraph.runs:
                    out.bold = True
                _shade(cell)
    _repeat_header(table.rows[0])
    doc.add_paragraph()


def _block(doc, block: Block) -> None:
    if block.kind == "paragraph":
        _runs(doc.add_paragraph(), block.runs)
    elif block.kind == "list":
        _list(doc, block)
    elif block.kind == "table":
        _table(doc, block)
    elif block.kind == "quote":
        _runs(doc.add_paragraph(style="Quote"), block.runs)
    elif block.kind == "code":
        out = doc.add_paragraph().add_run(_clean(block.text))
        out.font.name = "Consolas"
        out.font.size = Pt(9.5)
    elif block.kind == "heading":  # a heading nested inside a section's blocks
        _heading(doc, block.text, block.level)


# ---------------------------------------------------------------------------
# Page furniture: the disclosure on every page, and the page number
# ---------------------------------------------------------------------------


def _page_number(paragraph) -> None:
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        run = paragraph.add_run()
        if kind:
            field = OxmlElement("w:fldChar")
            field.set(qn("w:fldCharType"), kind)
            run._r.append(field)
        else:
            instruction = OxmlElement("w:instrText")
            instruction.set(qn("xml:space"), "preserve")
            instruction.text = text
            run._r.append(instruction)


def _small(paragraph, text: str) -> None:
    run = paragraph.add_run(text)
    run.font.size = Pt(8)
    run.font.color.rgb = _MUTED


def _furniture(doc, ir: ExportIRV1) -> None:
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _small(header, ir.disclosure_short)
    footer = section.footer.paragraphs[0]
    _small(footer, f"{ir.disclosure_long} · pagina ")
    _page_number(footer)


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------


def render_docx(ir: ExportIRV1, *, author: Optional[str] = None) -> DocxResult:
    """Render the snapshot to .docx bytes."""
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)

    props = doc.core_properties
    props.title = _clean(ir.title)
    props.author = author or "AixLearning — " + ir.disclosure_short
    props.subject = f"lesson={ir.lesson_id or 'n/a'} content={ir.content_hash[:16]}"
    props.comments = ir.disclosure_long
    props.category = f"aix-export {RENDERER_VERSION}"
    props.language = ir.language or "it"

    _furniture(doc, ir)
    doc.add_heading(_clean(ir.title), level=0)

    meta = []
    if ir.duration_minutes:
        meta.append(f"Durata: {ir.duration_minutes} minuti")
    meta.append(ir.disclosure_long)
    meta_paragraph = doc.add_paragraph()
    run = meta_paragraph.add_run(" · ".join(meta))
    run.font.color.rgb = _MUTED
    run.font.size = Pt(9.5)

    for block in ir.preamble:
        _block(doc, block)
    for section in ir.sections:
        _heading(doc, section.title, section.level)
        for block in section.blocks:
            _block(doc, block)

    buffer = io.BytesIO()
    doc.save(buffer)
    return DocxResult(content=buffer.getvalue())
