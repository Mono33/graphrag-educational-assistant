"""One parser per format, each turning bytes into ``DocumentIRV1`` (§7.6).

§5.4 fixes the order formats are switched on — digital PDF, DOCX, PPTX, then
images and scanned PDFs with OCR — and gives each its own flag, so a format that
misbehaves in the pilot can be switched off without switching off the others.
They are on by default: the whole pipeline is already behind
``AIX_MATERIALS_ENABLED``, so these flags only matter once that is on.

Every block records where it came from (``source_locator``) and how much to trust
it (``confidence``). A PDF page whose text layer yields fewer than
``min_text_chars`` characters is treated as having no text layer and handed to
OCR — which is the per-page decision the corpus showed was needed: the one real
scan in it has text on its cover page and nowhere else.

The parsers are synchronous: they take bytes and return an IR and touch no
database. The one network call is OCR with ``AIX_OCR_BACKEND=vision``, which
sends a page without text to a vision model. The worker runs them in a thread.
"""

from __future__ import annotations

import hashlib
import io
import logging
import os
import re
from collections.abc import Callable
from dataclasses import dataclass
from typing import Optional

from aix.webui.materials.parsing import intake
from aix.webui.materials.parsing.ir import (
    LOW_CONFIDENCE,
    DocumentIRV1,
    IRBlock,
    content_hash,
    ocr_index_min_confidence,
)
from aix.webui.materials.parsing.ocr import (
    DEFAULT_PAGE_TIMEOUT_S,
    ocr_available,
    recognise,
    render_pdf_page,
)

logger = logging.getLogger(__name__)

#: Recorded in ``ingestion_job.parser_profile_version`` and on every IR. A change
#: to what a parser extracts is a new profile, and a new profile is a re-parse.
PARSER_PROFILE = "ir@1"


@dataclass(frozen=True)
class ParseOptions:
    ocr_enabled: bool = True
    #: A PDF page with less extracted text than this has no *useful* text layer.
    min_text_chars: int = 25
    page_timeout_s: int = DEFAULT_PAGE_TIMEOUT_S
    max_pages: int = intake.MAX_PAGES
    #: OCR is the expensive part. Measured on the corpus scan under 2 vCPU / 4 GB:
    #: under 1 s a page at 200 dpi, ~55 MB for Tesseract. 30 pages keeps one upload
    #: to well under a minute of worker time. Pages beyond it stay unread, with a
    #: warning.
    max_ocr_pages: int = 30


def _flag(name: str) -> bool:
    return (os.getenv(name) or "true").strip().lower() not in ("0", "false", "no", "off")


class _Blocks:
    """Collects blocks with stable ids and a running reading order."""

    def __init__(self, sha: str) -> None:
        self._sha = sha
        self.items: list[IRBlock] = []

    def add(self, kind, text, locator, *, page=None, level=None, rows=(), confidence=1.0):
        text = (text or "").strip()
        if not text and not rows:
            return
        order = len(self.items)
        seed = f"{self._sha}:{order}:{kind}:{text[:200]}"
        self.items.append(
            IRBlock(
                block_id=hashlib.sha256(seed.encode()).hexdigest()[:16],
                kind=kind,
                text=text,
                reading_order=order,
                source_locator=locator,
                confidence=confidence,
                page=page,
                level=level,
                rows=tuple(tuple(cell for cell in row) for row in rows),
                content_hash=content_hash(text or repr(rows)),
            )
        )


def _paragraphs(text: str) -> list[str]:
    return [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def parse_pdf(content: bytes, options: ParseOptions) -> DocumentIRV1:
    from pypdf import PdfReader

    sha = hashlib.sha256(content).hexdigest()
    reader = PdfReader(io.BytesIO(content))
    if reader.is_encrypted:
        # Intake only lets through PDFs that open with an empty password.
        reader.decrypt("")
    total = len(reader.pages)
    blocks = _Blocks(sha)
    warnings: list[str] = []
    ocr_pages: list[int] = []
    unread: list[int] = []
    ocr_attempts = 0

    if total > options.max_pages:
        warnings.append(f"page_limit:{options.max_pages}")

    can_ocr = options.ocr_enabled and _flag("AIX_PARSE_OCR_ENABLED") and ocr_available()

    for index, page in enumerate(reader.pages[: options.max_pages]):
        number = index + 1
        text = page.extract_text() or ""

        if len(text.strip()) >= options.min_text_chars:
            for paragraph in _paragraphs(text):
                blocks.add("paragraph", paragraph, {"format": "pdf", "page": number}, page=number)
            continue

        if not can_ocr:
            unread.append(number)
            continue

        if ocr_attempts >= options.max_ocr_pages:
            if f"ocr_page_limit:{options.max_ocr_pages}" not in warnings:
                warnings.append(f"ocr_page_limit:{options.max_ocr_pages}")
            unread.append(number)
            continue
        ocr_attempts += 1

        try:
            result = recognise(render_pdf_page(content, index), timeout_s=options.page_timeout_s)
        except Exception:  # noqa: BLE001 — one unreadable page must not lose the rest
            logger.info("[parsing] OCR failed on page %d", number, exc_info=True)
            warnings.append(f"ocr_failed:page={number}")
            unread.append(number)
            continue

        if not result.text.strip():
            unread.append(number)
            continue

        # Read, but not trusted: kept out of the index and said so. See
        # ``OCR_INDEX_MIN_CONFIDENCE`` for where the floor comes from.
        if result.confidence < ocr_index_min_confidence():
            warnings.append(f"ocr_discarded:page={number}")
            continue

        ocr_pages.append(number)
        if result.confidence < LOW_CONFIDENCE:
            warnings.append(f"ocr_low_confidence:page={number}")
        for paragraph in _paragraphs(result.text) or [result.text]:
            blocks.add(
                "paragraph",
                paragraph,
                {"format": "pdf", "page": number, "ocr": True, "ocr_engine": result.engine},
                page=number,
                confidence=result.confidence,
            )

    if unread and not can_ocr:
        warnings.append("ocr_unavailable")

    return DocumentIRV1(
        parser_name="pdf",
        # 2: OCR pages under the confidence floor are not indexed.
        # 3: OCR may come from a vision model; the locator says which engine.
        parser_version="3",
        parser_profile_version=PARSER_PROFILE,
        source_sha256=sha,
        media_type=intake.PDF,
        blocks=tuple(blocks.items),
        page_count=total,
        ocr_pages=tuple(ocr_pages),
        unread_pages=tuple(unread),
        warnings=tuple(warnings),
    )


# ---------------------------------------------------------------------------
# Images inside Word and PowerPoint files
# ---------------------------------------------------------------------------

#: Shorter than this on either side: an icon, a bullet, a divider. Nothing to read.
MIN_EMBEDDED_SIDE_PX = 120

#: Fewer words than this once ``[figura]``/``[illeggibile]`` are removed: a photo
#: or a drawing without text. Not indexed, and not a warning either.
MIN_EMBEDDED_WORDS = 3


def _readable_words(text: str) -> int:
    return len(re.findall(r"\w{2,}", re.sub(r"\[[^\]]*\]", " ", text or "")))


class _EmbeddedImages:
    """OCR of the pictures inside a Word file or a deck, under the document's OCR budget.

    A slide that is a screenshot, a map or a chart holds its text in the picture,
    and it was skipped. Each distinct picture is read once (a logo repeated on
    every slide is read on the first), icons are skipped, and the same floor
    and page cap as scanned PDFs apply.
    """

    def __init__(self, options: ParseOptions) -> None:
        self.options = options
        self.enabled = options.ocr_enabled and _flag("AIX_PARSE_OCR_ENABLED") and ocr_available()
        self.seen: set[str] = set()
        self.attempts = 0
        self.pages: list[int] = []
        self.warnings: list[str] = []

    def read(self, blob: bytes, where: int):
        """The picture's text as an ``OcrResult``, or ``None`` when there is none to index."""
        if not self.enabled or not blob:
            return None
        digest = hashlib.sha1(blob).hexdigest()
        if digest in self.seen:
            return None
        self.seen.add(digest)

        from PIL import Image

        try:
            image = Image.open(io.BytesIO(blob))
            image.load()
        except Exception:  # noqa: BLE001 — WMF/EMF vectors and broken images carry no raster
            return None
        if min(image.size) < MIN_EMBEDDED_SIDE_PX:
            return None

        if self.attempts >= self.options.max_ocr_pages:
            code = f"ocr_page_limit:{self.options.max_ocr_pages}"
            if code not in self.warnings:
                self.warnings.append(code)
            return None
        self.attempts += 1

        try:
            result = recognise(image.convert("RGB"), timeout_s=self.options.page_timeout_s)
        except Exception:  # noqa: BLE001 — one unreadable picture must not lose the file
            logger.info("[parsing] OCR failed on an embedded image (%d)", where, exc_info=True)
            self.warnings.append(f"ocr_failed:page={where}")
            return None

        if _readable_words(result.text) < MIN_EMBEDDED_WORDS:
            return None
        if result.confidence < ocr_index_min_confidence():
            self.warnings.append(f"ocr_discarded:page={where}")
            return None
        if result.confidence < LOW_CONFIDENCE:
            self.warnings.append(f"ocr_low_confidence:page={where}")
        if where not in self.pages:
            self.pages.append(where)
        return result


def _add_image_text(blocks: _Blocks, result, locator: dict, *, page: Optional[int] = None) -> None:
    for paragraph in _paragraphs(result.text) or [result.text]:
        if not _readable_words(paragraph):
            continue  # a lone ``[figura]`` marks where a drawing was, not text
        blocks.add(
            "paragraph",
            paragraph,
            {**locator, "image": True, "ocr": True, "ocr_engine": result.engine},
            page=page,
            confidence=result.confidence,
        )


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def parse_docx(content: bytes, options: ParseOptions) -> DocumentIRV1:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    sha = hashlib.sha256(content).hexdigest()
    document = docx.Document(io.BytesIO(content))
    blocks = _Blocks(sha)
    unsupported: set[str] = set()
    images = _EmbeddedImages(options)
    image_number = 0

    # Walk the body in document order: paragraphs and tables interleave, and
    # python-docx's ``paragraphs`` and ``tables`` lists would lose that order.
    for index, element in enumerate(document.element.body.iterchildren()):
        tag = element.tag.rsplit("}", 1)[-1]
        locator = {"format": "docx", "body_index": index}

        if tag == "p":
            paragraph = Paragraph(element, document)
            style = (paragraph.style.name if paragraph.style is not None else "") or ""
            text = paragraph.text
            embedded = element.xpath(".//a:blip/@r:embed")
            if not embedded and element.xpath(".//w:pict"):
                unsupported.add("image")  # legacy VML pictures carry no blip to read
            heading = re.match(r"(?:Heading|Titolo)\s*(\d)", style, re.I)
            if heading:
                blocks.add("heading", text, locator, level=int(heading.group(1)))
            elif style.lower() == "title":
                blocks.add("heading", text, locator, level=1)
            elif "list" in style.lower() or element.xpath(".//w:numPr"):
                blocks.add("list_item", text, locator)
            else:
                blocks.add("paragraph", text, locator)

            # The pictures of this paragraph, read after its text, where they sit.
            for rel_id in embedded:
                part = document.part.related_parts.get(rel_id)
                blob = getattr(part, "blob", None)
                image_number += 1
                result = images.read(blob, image_number)
                if result is not None:
                    _add_image_text(blocks, result, {**locator, "image_index": image_number})

        elif tag == "tbl":
            table = Table(element, document)
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            blocks.add("table", " | ".join(" ".join(r) for r in rows), locator, rows=rows)

    return DocumentIRV1(
        parser_name="docx",
        # 2: the text inside embedded pictures is read.
        parser_version="2",
        parser_profile_version=PARSER_PROFILE,
        source_sha256=sha,
        media_type=intake.DOCX,
        blocks=tuple(blocks.items),
        warnings=tuple(images.warnings),
        unsupported=tuple(sorted(unsupported)),
    )


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------


def parse_pptx(content: bytes, options: ParseOptions) -> DocumentIRV1:
    from pptx import Presentation

    sha = hashlib.sha256(content).hexdigest()
    presentation = Presentation(io.BytesIO(content))
    blocks = _Blocks(sha)
    unsupported: set[str] = set()
    images = _EmbeddedImages(options)

    for slide_index, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        if title_shape is not None and title_shape.has_text_frame:
            blocks.add(
                "slide_title",
                title_shape.text_frame.text,
                {"format": "pptx", "slide_index": slide_index, "shape_id": title_shape.shape_id},
                page=slide_index,
            )

        # Grouped shapes are walked too: text and pictures inside a group were
        # skipped entirely.
        for shape in _leaf_shapes(slide.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            locator = {"format": "pptx", "slide_index": slide_index, "shape_id": shape.shape_id}

            if shape.has_table if hasattr(shape, "has_table") else False:
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                blocks.add(
                    "table",
                    " | ".join(" ".join(r) for r in rows),
                    locator,
                    page=slide_index,
                    rows=rows,
                )
            elif shape.has_text_frame:
                for paragraph_index, paragraph in enumerate(shape.text_frame.paragraphs):
                    text = "".join(run.text for run in paragraph.runs)
                    kind = "list_item" if paragraph.level > 0 else "paragraph"
                    blocks.add(
                        kind,
                        text,
                        {**locator, "paragraph_index": paragraph_index},
                        page=slide_index,
                    )
            elif _picture_blob(shape) is not None:
                result = images.read(_picture_blob(shape), slide_index)
                if result is not None:
                    _add_image_text(blocks, result, locator, page=slide_index)

        # The speaker notes. Teachers write the lesson there -- what to say,
        # what the slide only hints at -- and the first version of this parser
        # read the slides and skipped the notes entirely: on a deck built from
        # the ministerial curriculum, 4 of 4 sentences that lived only in the
        # notes were missing from the index (2026-09-25). They belong to their
        # slide, so a citation of them points at the slide they annotate.
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame
            for paragraph_index, paragraph in enumerate(notes.paragraphs):
                blocks.add(
                    "paragraph",
                    "".join(run.text for run in paragraph.runs),
                    {
                        "format": "pptx",
                        "slide_index": slide_index,
                        "notes": True,
                        "paragraph_index": paragraph_index,
                    },
                    page=slide_index,
                )

    return DocumentIRV1(
        parser_name="pptx",
        # 2: speaker notes are read. 3: pictures and grouped shapes are read.
        parser_version="3",
        parser_profile_version=PARSER_PROFILE,
        source_sha256=sha,
        media_type=intake.PPTX,
        blocks=tuple(blocks.items),
        page_count=len(presentation.slides),
        ocr_pages=tuple(images.pages),
        warnings=tuple(images.warnings),
        unsupported=tuple(sorted(unsupported)),
    )


def _leaf_shapes(shapes):
    """Every shape of a slide, descending into groups."""
    for shape in shapes:
        if getattr(shape, "shape_type", None) is not None and "GROUP" in str(shape.shape_type):
            yield from _leaf_shapes(shape.shapes)
        else:
            yield shape


def _picture_blob(shape) -> Optional[bytes]:
    """The bytes of a picture or picture placeholder; ``None`` for anything else
    (including a linked picture, whose bytes are not in the file)."""
    try:
        return shape.image.blob
    except Exception:  # noqa: BLE001 — not a picture, or not embedded
        return None


# ---------------------------------------------------------------------------
# Text and images
# ---------------------------------------------------------------------------


def parse_text(
    content: bytes, options: ParseOptions, media_type: str = intake.TEXT
) -> DocumentIRV1:
    sha = hashlib.sha256(content).hexdigest()
    blocks = _Blocks(sha)
    fmt = "markdown" if media_type == intake.MARKDOWN else "text"

    for index, paragraph in enumerate(_paragraphs(content.decode("utf-8"))):
        heading = re.match(r"^(#{1,6})\s+(.*)$", paragraph) if fmt == "markdown" else None
        locator = {"format": fmt, "paragraph_index": index}
        if heading and "\n" not in paragraph:
            blocks.add("heading", heading.group(2), locator, level=len(heading.group(1)))
        else:
            blocks.add("paragraph", paragraph, locator)

    return DocumentIRV1(
        parser_name=fmt,
        parser_version="1",
        parser_profile_version=PARSER_PROFILE,
        source_sha256=sha,
        media_type=media_type,
        blocks=tuple(blocks.items),
    )


def parse_image(content: bytes, options: ParseOptions, media_type: str) -> DocumentIRV1:
    from PIL import Image

    sha = hashlib.sha256(content).hexdigest()
    blocks = _Blocks(sha)
    warnings: list[str] = []
    ocr_pages: tuple[int, ...] = ()
    unread: tuple[int, ...] = ()

    if options.ocr_enabled and _flag("AIX_PARSE_OCR_ENABLED") and ocr_available():
        with Image.open(io.BytesIO(content)) as image:
            result = recognise(image, timeout_s=options.page_timeout_s)
        if result.text.strip() and result.confidence < ocr_index_min_confidence():
            warnings.append("ocr_discarded:page=1")
        elif result.text.strip():
            ocr_pages = (1,)
            if result.confidence < LOW_CONFIDENCE:
                warnings.append("ocr_low_confidence:page=1")
            for paragraph in _paragraphs(result.text) or [result.text]:
                blocks.add(
                    "paragraph",
                    paragraph,
                    {"format": "image", "page": 1, "ocr": True, "ocr_engine": result.engine},
                    page=1,
                    confidence=result.confidence,
                )
        else:
            unread = (1,)
    else:
        unread = (1,)
        warnings.append("ocr_unavailable")

    return DocumentIRV1(
        parser_name="image",
        # 2: an image read under the confidence floor is not indexed.
        # 3: OCR may come from a vision model; the locator says which engine.
        parser_version="3",
        parser_profile_version=PARSER_PROFILE,
        source_sha256=sha,
        media_type=media_type,
        blocks=tuple(blocks.items),
        page_count=1,
        ocr_pages=ocr_pages,
        unread_pages=unread,
        warnings=tuple(warnings),
    )


# ---------------------------------------------------------------------------
# Dispatch
# ---------------------------------------------------------------------------

_FORMAT_FLAGS = {
    intake.PDF: "AIX_PARSE_PDF_ENABLED",
    intake.DOCX: "AIX_PARSE_DOCX_ENABLED",
    intake.PPTX: "AIX_PARSE_PPTX_ENABLED",
}


class FormatDisabled(RuntimeError):
    """The format is accepted by intake but switched off for parsing."""


def parse_document(
    content: bytes, media_type: str, options: Optional[ParseOptions] = None
) -> DocumentIRV1:
    """Parse bytes that ``intake.check_upload`` has already accepted."""
    options = options or ParseOptions()

    flag = _FORMAT_FLAGS.get(media_type)
    if media_type in intake.IMAGE_TYPES:
        flag = "AIX_PARSE_IMAGES_ENABLED"
    if flag and not _flag(flag):
        raise FormatDisabled(f"{media_type} parsing is disabled by {flag}")

    parsers: dict[str, Callable[[], DocumentIRV1]] = {
        intake.PDF: lambda: parse_pdf(content, options),
        intake.DOCX: lambda: parse_docx(content, options),
        intake.PPTX: lambda: parse_pptx(content, options),
        intake.TEXT: lambda: parse_text(content, options, intake.TEXT),
        intake.MARKDOWN: lambda: parse_text(content, options, intake.MARKDOWN),
    }
    if media_type in intake.IMAGE_TYPES:
        return parse_image(content, options, media_type)
    if media_type not in parsers:
        raise ValueError(f"no parser for {media_type}")
    return parsers[media_type]()
